"""Snapshots, retention, reproduction and all-or-nothing publication from one snapshot."""
from __future__ import annotations

import gzip
import json
import os
import sqlite3
from contextlib import closing
from datetime import datetime, timedelta
from pathlib import Path

import pytest

import cyber_data_collector.quality.gate as gate
from cyber_data_collector.publication import outputs as pub
from cyber_data_collector.publication import snapshot as snaps
from cyber_data_collector.quality.model import load_thresholds


@pytest.fixture
def live(tmp_path):
    """A small WAL-mode database with committed changes still in its WAL."""
    path = tmp_path / "instance" / "live.db"
    path.parent.mkdir()
    conn = sqlite3.connect(path)
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("CREATE TABLE items (n INTEGER)")
    conn.executemany("INSERT INTO items VALUES (?)", [(i,) for i in range(5)])
    conn.commit()
    yield path, conn
    conn.close()


def _count(path) -> int:
    with closing(sqlite3.connect(path)) as c:
        return c.execute("SELECT COUNT(*) FROM items").fetchone()[0]


# --- snapshots --------------------------------------------------------------------------------

def test_snapshot_is_consistent_named_by_its_hash_and_leaves_the_source_alone(live, tmp_path):
    path, conn = live
    before = path.stat().st_mtime_ns
    snap = snaps.take(path, tmp_path / "work", now=datetime(2026, 10, 29, 18, 0, 0))
    assert _count(snap.path) == 5                                   # includes what is only in the WAL
    assert snap.snapshot_id == f"snap-20261029180000-{snap.sha256[:12]}"
    assert snaps.file_sha256(snap.path) == snap.sha256
    assert path.stat().st_mtime_ns == before
    conn.execute("INSERT INTO items VALUES (99)")
    conn.commit()
    assert _count(snap.path) == 5                                   # later writes are not in it
    with closing(sqlite3.connect(snap.path)) as c:
        assert c.execute("PRAGMA journal_mode").fetchone()[0] == "delete"   # one file; nothing beside it


def test_a_missing_database_is_an_error(tmp_path):
    with pytest.raises(snaps.SnapshotError):
        snaps.take(tmp_path / "nope.db", tmp_path / "work")


def test_retain_and_restore_round_trip_and_verify_the_hash(live, tmp_path):
    path, _ = live
    store = snaps.SnapshotStore(tmp_path / "store")
    snap = snaps.take(path, tmp_path / "work")
    manifest = store.retain(snap, status=snaps.PUBLISHED, quality_run_id="qr-1")
    assert manifest["sha256"] == snap.sha256 and manifest["compressed_bytes"] > 0
    restored = store.restore(snap.snapshot_id, tmp_path / "out")
    assert snaps.file_sha256(restored) == snap.sha256 and _count(restored) == 5
    assert store.resolve(snaps.LATEST_PUBLISHED) == snap.snapshot_id
    # A corrupted copy is refused, not used.
    with gzip.open(store.copy_path(snap.snapshot_id), "wb") as fh:
        fh.write(b"not the database")
    with pytest.raises(snaps.SnapshotError, match="corrupt"):
        store.restore(snap.snapshot_id, tmp_path / "out2")


def test_unknown_and_expired_snapshots_are_reported(live, tmp_path):
    path, _ = live
    store = snaps.SnapshotStore(tmp_path / "store")
    with pytest.raises(snaps.SnapshotError, match="unknown"):
        store.manifest("snap-nope")
    with pytest.raises(snaps.SnapshotError, match="no published snapshot"):
        store.resolve(snaps.LATEST_PUBLISHED)
    snap = snaps.take(path, tmp_path / "work")
    store.retain(snap, status=snaps.PUBLISHED)
    store.copy_path(snap.snapshot_id).unlink()
    with pytest.raises(snaps.SnapshotError, match="expired"):
        store.restore(snap.snapshot_id, tmp_path / "out")


def _fake_retained(store, n, status, hours_ago):
    sid = f"snap-2026102{n % 10}0000{n:02d}-{status[:3]}{n:09d}"
    taken = (datetime(2026, 10, 29) - timedelta(hours=hours_ago)).isoformat()
    store.root.mkdir(parents=True, exist_ok=True)
    store.copy_path(sid).write_bytes(b"x" * 100)
    store._write(sid, {"snapshot_id": sid, "sha256": "0" * 64, "taken_at": taken, "status": status})
    return sid


def test_retention_keeps_the_last_published_the_last_few_others_and_anything_pinned(tmp_path):
    store = snaps.SnapshotStore(tmp_path / "store")
    published = [_fake_retained(store, i, snaps.PUBLISHED, hours_ago=100 - i) for i in range(1, 9)]
    blocked = [_fake_retained(store, 10 + i, snaps.BLOCKED, hours_ago=50 - i) for i in range(4)]
    pinned = published[0]                                           # the oldest, named by an analysis
    expired = store.prune(pinned=[pinned])
    assert set(expired) == set(published[1:3]) | set(blocked[:2])    # 8 published: keep 5 newest + pinned
    assert all(store.copy_path(s).is_file() for s in published[3:] + [pinned] + blocked[2:])
    assert not any(store.copy_path(s).is_file() for s in expired)
    assert store.manifest(expired[0])["status"] == snaps.EXPIRED           # manifests outlive the copy
    assert store.manifest(expired[0])["expired_status"] in (snaps.PUBLISHED, snaps.BLOCKED)
    assert store.prune(pinned=[pinned]) == []                       # idempotent


def test_analysis_input_retains_names_and_reproduces_its_snapshot(live, tmp_path):
    path, conn = live
    store = snaps.SnapshotStore(tmp_path / "store")
    out = tmp_path / "analysis" / "run"
    with snaps.AnalysisInput.open(path, store=store) as first:
        assert _count(first.path) == 5 and not first.reused
        record = json.loads(first.record(out, analysis="test").read_text(encoding="utf-8"))
        working = first.path
    assert not working.exists()                                     # the working copy is removed ...
    assert store.copy_path(record["snapshot_id"]).is_file()         # ... the retained one is not
    assert record["snapshot_sha256"] == store.manifest(record["snapshot_id"])["sha256"]
    assert record["reproduce"] == f"--snapshot-id {record['snapshot_id']}"
    assert snaps.pinned_by_analyses(tmp_path / "analysis") == [record["snapshot_id"]]
    conn.execute("DELETE FROM items")                               # the live database moves on
    conn.commit()
    with snaps.AnalysisInput.open(path, snapshot_id=record["snapshot_id"], store=store) as again:
        assert again.reused and _count(again.path) == 5
        assert again.snapshot.sha256 == record["snapshot_sha256"]


# --- stamping -----------------------------------------------------------------------------------

def _snap(tmp_path):
    return snaps.Snapshot("snap-20261029180000-abcdef123456", "ab" * 32, "2026-10-29T18:00:00", "live.db",
                          tmp_path / "x.db", 1)


def test_html_outputs_record_the_snapshot_in_meta_and_footer(tmp_path):
    page = tmp_path / "p.html"
    page.write_text("<!doctype html><html><head><title>t</title></head><body><p>x</p></body></html>", encoding="utf-8")
    pub.stamp_html(page, _snap(tmp_path), "2026-10-29T18:05:00")
    text = page.read_text(encoding="utf-8")
    assert '<meta name="data-snapshot" content="snap-20261029180000-abcdef123456">' in text
    assert text.index("data-snapshot") < text.index("<title>")
    assert "Data snapshot <code>snap-20261029180000-abcdef123456</code>" in text
    assert text.index("class=\"data-snapshot\"") < text.index("</body>")
    bare = tmp_path / "b.html"
    bare.write_text("<p>fragment</p>", encoding="utf-8")
    pub.stamp_html(bare, _snap(tmp_path), "now")
    assert "data-snapshot" in bare.read_text(encoding="utf-8")


def test_workbooks_record_the_snapshot_in_their_properties(tmp_path):
    from openpyxl import Workbook, load_workbook

    path = tmp_path / "m.xlsx"
    wb = Workbook()
    wb.active["A1"] = "keep me"
    wb.save(path)
    pub.stamp_workbook(path, _snap(tmp_path), "2026-10-29T18:05:00")
    loaded = load_workbook(path)
    assert loaded.active["A1"].value == "keep me"
    assert "snap-20261029180000-abcdef123456" in loaded.properties.description
    assert loaded.properties.keywords == "snapshot:snap-20261029180000-abcdef123456"


# --- install ---------------------------------------------------------------------------------------

def test_install_moves_everything_or_restores_everything(tmp_path):
    work = tmp_path / "work"
    work.mkdir()
    out = tmp_path / "dash"
    out.mkdir()
    (out / "a.html").write_text("old a")
    (out / "b.html").write_text("old b")
    staged = []
    for name in ("a.html", "b.html", "c.html"):
        f = work / f"new-{name}"
        f.write_text(f"new {name}")
        staged.append((f, out / name))
    staged.append((work / "missing-file", out / "d.html"))           # the fourth move fails
    with pytest.raises(pub.PublicationError, match="restored"):
        pub.install(staged, work / "journal-dir")
    assert (out / "a.html").read_text() == "old a" and (out / "b.html").read_text() == "old b"
    assert not (out / "c.html").exists() and not (out / "d.html").exists()
    # And the whole set goes in when nothing fails (the staged files went with the failed attempt).
    for source, _ in staged[:3]:
        source.write_text(f"new {source.name[4:]}")
    pub.install(staged[:3], work / "journal-dir2")
    assert [(out / n).read_text() for n in ("a.html", "b.html", "c.html")] == [
        "new a.html", "new b.html", "new c.html"]


def test_an_interrupted_install_is_rolled_back_by_the_next_publication(tmp_path):
    out = tmp_path / "dash"
    out.mkdir()
    (out / "a.html").write_text("old a")
    work = tmp_path / ".publish-crashed"
    (work / "backup").mkdir(parents=True)
    (work / "backup" / "0").write_text("old a")
    (out / "a.html").write_text("half-published a")
    (out / "new.html").write_text("half-published new")
    (work / pub.JOURNAL).write_text(json.dumps({"complete": False, "entries": [
        {"target": str(out / "a.html"), "backup": str(work / "backup" / "0")},
        {"target": str(out / "new.html"), "backup": None}]}))
    assert pub.recover_interrupted(tmp_path) == [".publish-crashed"]
    assert (out / "a.html").read_text() == "old a" and not (out / "new.html").exists()
    assert not work.exists()


# --- publish -----------------------------------------------------------------------------------------

def _report(snap, allowed=True, reasons=()):
    return gate.QualityReport(run_id="qr-test-1", created_at="2026-10-29T18:00:00", db_path="live",
                              thresholds=load_thresholds(), metrics=[], allowed=allowed, reasons=list(reasons),
                              snapshot_id=snap.snapshot_id, snapshot_sha256=snap.sha256)


@pytest.fixture
def harness(live, tmp_path, monkeypatch):
    """Publication with fake builders and a fake gate; records which snapshot each stage read."""
    path, conn = live
    seen = {"gate": [], "build": [], "recorded": []}
    verdict = {"allowed": True, "reasons": [], "mutate_during_gate": False}

    def fake_checks(c, thresholds=None, today=None, evaluation_dir=None, snapshot=None, db_path=None):
        seen["gate"].append((snapshot.snapshot_id, c.execute("SELECT COUNT(*) FROM items").fetchone()[0]))
        if verdict["mutate_during_gate"]:                       # the live database moves while we publish
            conn.execute("INSERT INTO items VALUES (100)")
            conn.commit()
        return _report(snapshot, verdict["allowed"], verdict["reasons"])

    def fake_store(c, report):
        seen["recorded"].append((report.run_id, report.snapshot_id, c.execute("PRAGMA database_list").fetchone()[2]))
        return report.run_id

    monkeypatch.setattr(gate, "run_quality_checks", fake_checks)
    monkeypatch.setattr(gate, "store_report", fake_store)

    def dashboard(ctx):
        n = _count(ctx.db)
        seen["build"].append((ctx.snapshot.snapshot_id, n))
        index = ctx.dashboard_dir / "index.html"
        index.write_text(f"<html><head></head><body>rows={n}</body></html>", encoding="utf-8")
        entities = ctx.dashboard_dir / "entities.html"
        entities.write_text(f"<html><body>entities rows={n}</body></html>", encoding="utf-8")
        return [index, entities]

    def risk(ctx):
        from openpyxl import Workbook
        wb = Workbook()
        wb.active["A1"] = _count(ctx.db)
        target = ctx.risk_dir / "risk_matrix.xlsx"
        wb.save(target)
        return [target]

    builders = {"dashboard": dashboard, "risk_matrix": risk}
    out, risk_dir, analysis = tmp_path / "dashboard", tmp_path / "risk_matrix", tmp_path / "analysis"
    store = snaps.SnapshotStore(tmp_path / "instance" / "snapshots")

    def run(**kw):
        return pub.publish(path, out, risk_dir, analysis_dir=analysis, builders=kw.pop("builders", builders),
                           store=store, repo_root=tmp_path, **kw)

    return type("H", (), dict(path=path, conn=conn, seen=seen, verdict=verdict, run=staticmethod(run),
                              out=out, risk=risk_dir, analysis=analysis, store=store, builders=builders))


def test_everything_is_built_from_the_one_snapshot_the_gate_read(harness):
    harness.verdict["mutate_during_gate"] = True               # live gets a 6th row after the snapshot
    result = harness.run()
    gate_id, gate_rows = harness.seen["gate"][0]
    build_id, build_rows = harness.seen["build"][0]
    assert gate_id == build_id == result.snapshot_id
    assert gate_rows == build_rows == 5 and _count(harness.path) == 6      # the outputs show 5, not 6
    assert "rows=5" in (harness.out / "index.html").read_text(encoding="utf-8")
    assert "entities rows=5" in (harness.out / "entities.html").read_text(encoding="utf-8")
    from openpyxl import load_workbook
    assert load_workbook(harness.risk / "risk_matrix.xlsx").active["A1"].value == 5


def test_every_output_names_the_snapshot_and_the_manifest_hashes_them(harness):
    result = harness.run()
    sid = result.snapshot_id
    for name in ("index.html", "entities.html"):
        assert f'content="{sid}"' in (harness.out / name).read_text(encoding="utf-8")
    from openpyxl import load_workbook
    assert sid in load_workbook(harness.risk / "risk_matrix.xlsx").properties.description
    manifest = json.loads((harness.out / "publication.json").read_text(encoding="utf-8"))
    assert manifest["snapshot_id"] == sid and manifest["quality_run_id"] == "qr-test-1"
    assert manifest["quality_gate"] == "passed" and manifest["snapshot_sha256"] == result.snapshot_sha256
    by_path = {Path(o["path"]).name: o for o in manifest["outputs"]}
    assert set(by_path) == {"index.html", "entities.html", "risk_matrix.xlsx"}
    assert by_path["index.html"]["sha256"] == snaps.file_sha256(harness.out / "index.html")
    assert by_path["risk_matrix.xlsx"]["sha256"] == snaps.file_sha256(harness.risk / "risk_matrix.xlsx")
    assert "scripts/publish_outputs.py --snapshot-id " + sid in manifest["reproduce"]
    assert harness.seen["recorded"][0][:2] == ("qr-test-1", sid)
    assert harness.seen["recorded"][0][2].endswith("live.db")             # recorded in the live database


def test_the_snapshot_is_retained_and_marked_published(harness):
    result = harness.run()
    manifest = harness.store.manifest(result.snapshot_id)
    assert manifest["status"] == snaps.PUBLISHED and manifest["quality_run_id"] == "qr-test-1"
    assert sorted(Path(p).name for p in manifest["outputs"]) == ["entities.html", "index.html", "risk_matrix.xlsx"]
    restored = harness.store.restore(result.snapshot_id, harness.out.parent / "restore")
    assert _count(restored) == 5 and snaps.file_sha256(restored) == result.snapshot_sha256
    assert not list(harness.out.parent.glob(".publish-*"))                 # staging removed


def test_a_blocked_gate_replaces_nothing_and_keeps_the_snapshot(harness):
    harness.run()                                                        # a good publication first
    before = {p: p.read_bytes() for p in (harness.out / "index.html", harness.out / "entities.html",
                                          harness.out / "publication.json", harness.risk / "risk_matrix.xlsx")}
    harness.conn.execute("INSERT INTO items VALUES (7)")
    harness.conn.commit()
    harness.verdict.update(allowed=False, reasons=["integrity.reporting-population: fail"])
    with pytest.raises(pub.PublicationBlocked, match="integrity.reporting-population") as raised:
        harness.run()
    assert {p: p.read_bytes() for p in before} == before
    blocked = harness.store.manifest(raised.value.snapshot_id)
    assert blocked["status"] == snaps.BLOCKED and blocked["reasons"] == ["integrity.reporting-population: fail"]
    assert len(harness.seen["build"]) == 1                                # nothing was built for the blocked run
    assert harness.seen["recorded"][-1][1] == raised.value.snapshot_id    # but the run (and its snapshot) is recorded


def test_a_failing_builder_publishes_nothing(harness):
    harness.run()
    old = (harness.out / "index.html").read_bytes()

    def boom(ctx):
        raise RuntimeError("query failed")

    builders = {**harness.builders, "entities": boom}
    harness.conn.execute("INSERT INTO items VALUES (8)")
    harness.conn.commit()
    with pytest.raises(pub.PublicationError, match="entities could not be built"):
        harness.run(builders=builders)
    assert (harness.out / "index.html").read_bytes() == old               # not even the page that built fine
    failed = [m for m in harness.store.manifests() if m["status"] == snaps.FAILED]
    assert len(failed) == 1


def test_a_failed_move_restores_the_previous_outputs(harness, monkeypatch):
    harness.run()
    old = {p: p.read_bytes() for p in (harness.out / "index.html", harness.risk / "risk_matrix.xlsx")}
    real_replace = os.replace
    calls = {"n": 0}

    def flaky(src, dst):
        if str(dst).endswith("risk_matrix.xlsx") and "backup" not in str(src):
            calls["n"] += 1
            raise OSError("disk full")
        return real_replace(src, dst)

    harness.conn.execute("INSERT INTO items VALUES (9)")
    harness.conn.commit()
    monkeypatch.setattr(pub.os, "replace", flaky)
    with pytest.raises(pub.PublicationError, match="restored"):
        harness.run()
    monkeypatch.undo()
    assert calls["n"] == 1
    assert {p: p.read_bytes() for p in old} == old
    assert "rows=5" in (harness.out / "index.html").read_text(encoding="utf-8")


def test_a_snapshot_is_reproduced_exactly_from_the_retained_copy(harness):
    first = harness.run()
    first_index = (harness.out / "index.html").read_text(encoding="utf-8")
    harness.conn.execute("INSERT INTO items VALUES (50)")                  # the live database has moved on
    harness.conn.commit()
    again = harness.run(snapshot_id=first.snapshot_id)
    assert again.snapshot_id == first.snapshot_id and again.snapshot_sha256 == first.snapshot_sha256
    body = lambda text: text.split("Data snapshot")[0].split("<head>")[-1]   # noqa: E731 - strip the timestamp
    assert body((harness.out / "index.html").read_text(encoding="utf-8")) == body(first_index)
    assert "rows=5" in (harness.out / "index.html").read_text(encoding="utf-8")
    assert harness.run(snapshot_id=snaps.LATEST_PUBLISHED).snapshot_id == first.snapshot_id


def test_skipping_the_gate_is_recorded_not_hidden(harness):
    result = harness.run(skip_gate=True)
    assert result.gate.startswith("skipped") and not harness.seen["gate"]
    manifest = json.loads((harness.out / "publication.json").read_text(encoding="utf-8"))
    assert manifest["quality_gate"].startswith("skipped") and manifest["quality_run_id"] is None


def test_analyses_on_another_snapshot_are_flagged_in_the_manifest(harness):
    run = harness.analysis / "recurrent_timing"
    run.mkdir(parents=True)
    (run / "analysis_inputs.json").write_text(json.dumps({"snapshot_id": "snap-old"}), encoding="utf-8")
    bare = harness.analysis / "recurrent_timing_90d"
    bare.mkdir()
    result = harness.run()
    assert any("recurrent_timing was computed on snapshot snap-old" in w for w in result.warnings)
    assert any("recurrent_timing_90d does not record" in w for w in result.warnings)
    manifest = json.loads((harness.out / "publication.json").read_text(encoding="utf-8"))
    assert manifest["analysis_snapshots"] == {"recurrent_timing": "snap-old", "recurrent_timing_90d": None}


def test_retention_runs_after_publication_and_never_expires_a_pinned_analysis_snapshot(harness):
    pinned = _fake_retained(harness.store, 1, snaps.PUBLISHED, hours_ago=999)
    run = harness.analysis / "recurrent_timing"
    run.mkdir(parents=True)
    (run / "analysis_inputs.json").write_text(json.dumps({"snapshot_id": pinned}), encoding="utf-8")
    for i in range(2, 8):
        _fake_retained(harness.store, i, snaps.PUBLISHED, hours_ago=500 - i)
    result = harness.run()
    assert harness.store.copy_path(pinned).is_file() and harness.store.copy_path(result.snapshot_id).is_file()
    published_kept = [m for m in harness.store.manifests()
                      if m["status"] == snaps.PUBLISHED and harness.store.copy_path(m["snapshot_id"]).is_file()]
    assert len(published_kept) == snaps.KEEP_PUBLISHED + 1                 # the five newest, and the pinned one


def test_only_one_publication_runs_at_a_time(harness):
    lock = harness.out.parent / pub.LOCK
    lock.write_text("someone else")
    with pytest.raises(pub.PublicationError, match="another publication is running"):
        harness.run()
    assert not (harness.out / "index.html").exists()
    old = (datetime.now() - timedelta(hours=pub.LOCK_STALE_HOURS + 1)).timestamp()
    os.utime(lock, (old, old))                                         # a crashed run's lock is taken over
    harness.run()
    assert (harness.out / "index.html").exists() and not lock.exists()


def test_default_builders_cover_every_published_output():
    assert set(pub.DEFAULT_BUILDERS) == {"dashboard", "entities", "timelines", "dedup", "recurrence", "risk_matrix"}


# --- quality runs name their snapshot ---------------------------------------------------------------------

def test_quality_runs_record_and_report_the_snapshot(tmp_path, live):
    path, _ = live
    snap = snaps.take(path, tmp_path / "work")
    conn = sqlite3.connect(tmp_path / "q.db")
    # An older database: the table exists without the snapshot columns; storing a run adds them.
    conn.execute("CREATE TABLE QualityRuns (run_id TEXT PRIMARY KEY, created_at TEXT NOT NULL, db_path TEXT, "
                 "thresholds_version TEXT NOT NULL, thresholds_sha256 TEXT NOT NULL, publication_allowed INTEGER "
                 "NOT NULL, blocking_failures INTEGER NOT NULL, failed INTEGER NOT NULL, warned INTEGER NOT NULL, "
                 "not_measured INTEGER NOT NULL, reasons_json TEXT NOT NULL)")
    report = _report(snap)
    gate.store_report(conn, report)
    assert conn.execute("SELECT run_id, snapshot_id, snapshot_sha256 FROM QualityRuns").fetchone() == (
        "qr-test-1", snap.snapshot_id, snap.sha256)
    text = gate.render_markdown(report)
    assert f"Snapshot: `{snap.snapshot_id}`" in text and f"--snapshot-id {snap.snapshot_id}" in text
    assert snap.snapshot_id in gate.render_html(report)
    assert report.to_dict()["snapshot_id"] == snap.snapshot_id
    live_report = gate.QualityReport("qr-2", "now", "live", load_thresholds(), [], True, [])
    assert "none - measured the live database directly" in gate.render_markdown(live_report)
    conn.close()


# --- recurrence page names the data -----------------------------------------------------------------------

def test_recurrence_page_warns_when_runs_use_different_or_unrecorded_snapshots(tmp_path):
    from scripts import build_recurrence_dashboard as brd

    for folder, sid in (("recurrent_timing", "snap-a"), ("recurrent_timing_90d", "snap-b"),
                        ("recurrent_timing_90_180_peak", None), (brd.EXTENDED_DIR, "snap-a")):
        d = tmp_path / folder
        d.mkdir()
        if sid:
            (d / "analysis_inputs.json").write_text(json.dumps({"snapshot_id": sid}), encoding="utf-8")
    notes = brd.snapshot_notes(tmp_path, "snap-a")
    assert any("recurrent_timing_90_180_peak" in n and "not recorded" in n for n in notes)
    assert any("different snapshots (snap-a, snap-b)" in n for n in notes)
    assert any("recurrent_timing_90d not computed on the published snapshot snap-a" in n for n in notes)
    assert brd.snapshot_notes(tmp_path / "empty", "snap-a") == []


# --- the extended analyses: one location, and an input that can be restored ------------------------------

def _fake_extended(monkeypatch):
    import pandas as pd
    from datetime import date

    from scripts import analyze_recurrence_extensions as runner

    monkeypatch.setattr(runner, "load_victim_events", lambda _: pd.DataFrame([
        {"deduplicated_event_id": "e1", "event_date": date(2024, 1, 1)}]))
    monkeypatch.setattr(runner, "add_first_discovery", lambda events, _: events.assign(first_discovered_at="2024-01-03"))
    seen = []

    def fake_run(events, snapshot, end, out_dir, seed, **kwargs):
        seen.append(snapshot)
        with closing(sqlite3.connect(snapshot)) as c:
            return {"section": "## S", "rows": c.execute("SELECT COUNT(*) FROM items").fetchone()[0]}

    for module in (runner.timing_precision, runner.linkage_supplier, runner.sequence_bias):
        monkeypatch.setattr(module, "run", fake_run)
    return runner, seen


def test_extended_analysis_retains_its_snapshot_and_names_it_beside_the_outputs(live, tmp_path, monkeypatch):
    from datetime import date

    path, conn = live
    runner, seen = _fake_extended(monkeypatch)
    out = tmp_path / "analysis" / "recurrent_timing_extended"
    summary = runner.run_analysis(path, out, date(2024, 3, 1), bootstrap=2, date_draws=2)
    record = json.loads((out / snaps.ANALYSIS_INPUTS).read_text(encoding="utf-8"))
    store = snaps.store_for(path)
    assert record["snapshot_id"] == summary["snapshot_id"] and store.copy_path(record["snapshot_id"]).is_file()
    assert store.manifest(record["snapshot_id"])["status"] == snaps.ANALYSIS
    assert "retained under `instance/snapshots/`" in (out / "extended_recurrence_report.md").read_text(encoding="utf-8")
    assert not seen[0].exists()                                      # the working copy is gone; the retained one is not
    conn.execute("INSERT INTO items VALUES (77)")                    # the live database moves on
    conn.commit()
    again = runner.run_analysis(path, tmp_path / "again", date(2024, 3, 1), bootstrap=2, date_draws=2,
                                snapshot_id=record["snapshot_id"])
    assert again["snapshot_sha256"] == summary["snapshot_sha256"]
    assert again["analyses"]["timing_and_date_precision"]["rows"] == 5


def test_analyses_sharing_a_snapshot_take_only_one(live, tmp_path, monkeypatch):
    from datetime import date

    path, _ = live
    runner, _ = _fake_extended(monkeypatch)
    with snaps.AnalysisInput.open(path) as shared:
        runner.run_analysis(path, tmp_path / "a", date(2024, 3, 1), bootstrap=2, date_draws=2, analysis_input=shared)
        runner.run_analysis(path, tmp_path / "b", date(2024, 3, 1), bootstrap=2, date_draws=2, analysis_input=shared)
        assert shared.path.exists()                                  # the caller's snapshot is not closed by a run
    assert len(snaps.store_for(path).manifests()) == 1
    ids = {json.loads((tmp_path / d / snaps.ANALYSIS_INPUTS).read_text(encoding="utf-8"))["snapshot_id"]
           for d in ("a", "b")}
    assert len(ids) == 1


def test_the_repository_has_one_extended_results_folder_with_the_frozen_cohorts():
    from scripts import analyze_recurrence_extensions as runner

    repo = Path(__file__).resolve().parents[2]
    assert runner.DEFAULT_OUT == "analysis/recurrent_timing_extended"
    assert not (repo / "analysis" / "recurrent_timing" / "extended").exists()
    for frozen in ("timing_peak_registration.json", "timing_size_adjusted_registration.json"):
        assert (repo / "analysis" / "recurrent_timing_extended" / frozen).is_file()
